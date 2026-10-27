from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from copy import copy
from openpyxl.styles import PatternFill, Font


# ==================================================
# 1. CONFIGURATION
# ==================================================

PROJECT_DIR = Path(__file__).resolve().parent.parent

REPORTS_DIR = PROJECT_DIR / "reports"
CHARTS_DIR = PROJECT_DIR / "outputs" / "charts"

OUTPUT_FILE = REPORTS_DIR / "rapport_commercial.xlsx"


# ==================================================
# 2. LECTURE DES RAPPORTS CSV
# ==================================================

reports = {
    "Synthèse KPI": "kpi_summary.csv",
    "Ventes par catégorie": "sales_by_category.csv",
    "Ventes par région": "sales_by_region.csv",
    "Top 10 produits": "top_10_products.csv",
    "Ventes mensuelles": "monthly_sales.csv",
}

dataframes = {}

for sheet_name, filename in reports.items():
    file_path = REPORTS_DIR / filename

    if not file_path.exists():
        raise FileNotFoundError(
            f"Fichier introuvable : {file_path}"
        )

    dataframes[sheet_name] = pd.read_csv(file_path)

print("Rapports CSV chargés avec succès.")


# ==================================================
# 3. EXPORT VERS EXCEL
# ==================================================

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    for sheet_name, dataframe in dataframes.items():
        dataframe.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )

print("Données exportées vers Excel.")


# ==================================================
# 4. MISE EN FORME DU CLASSEUR
# ==================================================

workbook = load_workbook(OUTPUT_FILE)

for worksheet in workbook.worksheets:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Mise en forme des en-têtes
    for cell in worksheet[1]:
        new_font = copy(cell.font)
        new_font.bold = True
        new_font.color = "FFFFFF"
        cell.font = new_font

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

    # Ajustement automatique des largeurs
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 35)


# ==================================================
# 5. AJOUT DES GRAPHIQUES
# ==================================================

charts = [
    (
        "monthly_revenue.png",
        "A10"
    ),
    (
        "revenue_by_category.png",
        "J10"
    ),
    (
        "revenue_by_region.png",
        "A35"
    ),
    (
        "top_10_products.png",
        "J35"
    ),
]

dashboard = workbook.create_sheet("Dashboard")

dashboard["A1"] = "SUNU MARKET — RAPPORT COMMERCIAL"

dashboard["A1"].font = Font(
    bold=True,
    size=18,
    color="1F4E78"
)

dashboard["A3"] = (
    "Analyse automatisée des ventes — données fictives"
)

dashboard["A3"].font = Font(
    italic=True,
    color="666666"
)

for filename, position in charts:
    chart_path = CHARTS_DIR / filename

    if not chart_path.exists():
        print(
            f"Attention : graphique introuvable : {filename}"
        )
        continue

    image = Image(str(chart_path))

    image.width = 600
    image.height = 350

    dashboard.add_image(image, position)

dashboard.sheet_view.showGridLines = False

workbook.save(OUTPUT_FILE)


# ==================================================
# 6. RÉSULTAT
# ==================================================

print("=" * 50)
print("RAPPORT COMMERCIAL GÉNÉRÉ AVEC SUCCÈS")
print("=" * 50)
print(f"Fichier créé : {OUTPUT_FILE}")
print(f"Nombre de feuilles : {len(workbook.sheetnames)}")
print("Feuilles :", ", ".join(workbook.sheetnames))